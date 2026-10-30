"""The master list. master.json holds every employee (standard schema); verified ones are written to the
master Excel sheet. Writes are atomic (temp file + rename) and a .bak of the previous version is kept."""
import json
import os
import shutil
import threading
from pathlib import Path

from .schema import now_iso
from .validate import normalize_phone

COLUMNS = [
    ("Employee ID", lambda e: e["employee_id"]), ("Full Name", lambda e: e["full_name"]), ("Email", lambda e: e["email"]),
    ("Phone", lambda e: e["phone"]), ("Age", lambda e: e["age"] if e["age"] is not None else ""), ("Date of Birth", lambda e: e["dob"]),
    ("Gender", lambda e: e["gender"]), ("Father's Name", lambda e: e["father_name"]), ("Address", lambda e: e["address"]["full"]),
    ("Pincode", lambda e: e["address"]["pincode"]), ("Aadhaar (masked)", lambda e: e["aadhaar"]["number_masked"]),
    ("PAN", lambda e: e["pan"]["number"]), ("Account Holder", lambda e: e["bank"]["account_holder"]),
    ("Account Number", lambda e: e["bank"]["account_number"]), ("IFSC", lambda e: e["bank"]["ifsc"]),
    ("Bank", lambda e: e["bank"]["bank_name"]), ("Branch", lambda e: e["bank"]["branch"]),
    ("Other Document", lambda e: e["other"]["reference"]),
    ("Data", lambda e: "placeholder" if any(d.get("source") == "placeholder" for d in e["documents"].values()) else "entered"),
    ("Verified At", lambda e: e["verification"]["checked_at"]),
]


class Duplicate(Exception):
    def __init__(self, conflicts: list[dict]):
        self.conflicts = conflicts
        super().__init__("; ".join(f"{c['field']} already registered to {c['employee_id']}" for c in conflicts))


def identity_keys(e: dict) -> dict[str, str]:
    keys = {
        "email": (e.get("email") or "").strip().lower(),
        "phone": normalize_phone(e.get("phone", "")),
        "PAN": (e["pan"].get("number") or "").upper(),
        "Aadhaar": e["aadhaar"].get("hash") or "",
        "bank account": f"{(e['bank'].get('ifsc') or '')[:4]}:{e['bank'].get('account_number')}" if e["bank"].get("account_number") else "",
    }
    return {k: v for k, v in keys.items() if v}


class MasterStore:
    def __init__(self, data_dir: str | Path):
        self.dir = Path(data_dir)
        self.dir.mkdir(parents=True, exist_ok=True)
        self.json_path = self.dir / "master.json"
        self.xlsx_path = self.dir / "master_employees.xlsx"
        self.xls_path = self.dir / "master_employees.xls"
        self._lock = threading.RLock()
        if not self.json_path.exists():
            self._write({"version": 1, "employees": []})

    # ---- persistence ----
    def _read(self) -> dict:
        try:
            return json.loads(self.json_path.read_text("utf-8"))
        except (json.JSONDecodeError, OSError):
            bak = self.json_path.with_suffix(".json.bak")
            if bak.exists():
                return json.loads(bak.read_text("utf-8"))
            raise RuntimeError("master.json is unreadable and there is no backup. Fix or restore it before continuing.")

    def _write(self, data: dict) -> None:
        tmp = self.json_path.with_suffix(".json.tmp")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
            f.flush()
            os.fsync(f.fileno())
        if self.json_path.exists():
            shutil.copyfile(self.json_path, self.json_path.with_suffix(".json.bak"))
        os.replace(tmp, self.json_path)

    # ---- queries ----
    def all(self) -> list[dict]:
        with self._lock:
            return self._read()["employees"]

    def get(self, emp_id: str) -> dict | None:
        return next((e for e in self.all() if e["employee_id"] == emp_id), None)

    def find_conflicts(self, emp: dict, exclude_id: str | None = None) -> list[dict]:
        mine = identity_keys(emp)
        out = []
        for e in self.all():
            if e["employee_id"] == exclude_id:
                continue
            theirs = identity_keys(e)
            for field, val in mine.items():
                if theirs.get(field) == val:
                    out.append({"field": field, "employee_id": e["employee_id"], "name": e["full_name"]})
        return out

    # ---- writes ----
    def add(self, emp: dict) -> dict:
        with self._lock:
            data = self._read()
            conflicts = self.find_conflicts(emp)
            if conflicts:
                raise Duplicate(conflicts)
            n = max([int(e["employee_id"].split("-")[1]) for e in data["employees"]] + [0]) + 1
            emp["employee_id"] = f"EMP-{n:04d}"
            emp["status"] = "COMPLETE"
            emp["updated_at"] = now_iso()
            data["employees"].append(emp)
            self._write(data)
            return emp

    def update(self, emp: dict, publish: bool = True) -> dict:
        with self._lock:
            data = self._read()
            conflicts = self.find_conflicts(emp, exclude_id=emp["employee_id"])
            if conflicts:
                raise Duplicate(conflicts)
            idx = next((i for i, e in enumerate(data["employees"]) if e["employee_id"] == emp["employee_id"]), None)
            if idx is None:
                raise KeyError(emp["employee_id"])
            was_verified = data["employees"][idx]["status"] == "VERIFIED"
            emp["updated_at"] = now_iso()
            data["employees"][idx] = emp
            self._write(data)
            if publish and (was_verified or emp["status"] == "VERIFIED"):
                self.publish_excel(data["employees"])
            return emp

    # ---- master sheet ----
    def publish_excel(self, employees: list[dict] | None = None) -> int:
        """Rebuild the master sheet from the VERIFIED employees. Rebuilding (not appending) makes duplicates impossible."""
        with self._lock:
            rows = [e for e in (employees if employees is not None else self.all()) if e["status"] == "VERIFIED"]
            rows.sort(key=lambda e: e["employee_id"])
            self._write_xlsx(rows)
            self._write_xls(rows)
            return len(rows)

    def _write_xlsx(self, rows: list[dict]) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Employees"
        ws.append([c for c, _ in COLUMNS])
        for e in rows:
            ws.append([str(fn(e)) if isinstance(fn(e), str) else fn(e) for _, fn in COLUMNS])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1D1D1F")
        for i, (name, _) in enumerate(COLUMNS, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, min(48, len(name) + 6))
        ws.freeze_panes = "C2"
        tmp = self.xlsx_path.with_suffix(".xlsx.tmp")
        wb.save(tmp)
        os.replace(tmp, self.xlsx_path)

    def _write_xls(self, rows: list[dict]) -> None:
        import xlwt
        wb = xlwt.Workbook(encoding="utf-8")
        ws = wb.add_sheet("Employees")
        bold = xlwt.easyxf("font: bold on")
        for j, (name, _) in enumerate(COLUMNS):
            ws.write(0, j, name, bold)
            ws.col(j).width = 256 * max(14, min(48, len(name) + 6))
        for i, e in enumerate(rows, 1):
            for j, (_, fn) in enumerate(COLUMNS):
                v = fn(e)
                ws.write(i, j, v if isinstance(v, (int, float)) else str(v))
        tmp = self.xls_path.with_suffix(".xls.tmp")
        wb.save(str(tmp))
        os.replace(tmp, self.xls_path)
