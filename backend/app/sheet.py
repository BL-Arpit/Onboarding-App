"""Reads the HR sheet: three columns (name, email, phone) as .xlsx, .xls or .csv. Header is optional."""
import csv
import io
import re

from rapidfuzz import fuzz, process

from .validate import normalize_phone, valid_email

SYN = {
    "name": ["name", "fullname", "employeename", "empname", "nameofemployee", "employee"],
    "email": ["email", "emailid", "emailaddress", "mail", "mailid"],
    "phone": ["phone", "phoneno", "phonenumber", "mobile", "mobileno", "mobilenumber", "contact", "contactno", "cell"],
}
_FLAT = {s: k for k, v in SYN.items() for s in v}


class SheetError(Exception):
    pass


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))          # Excel stores 9845012345 as 9845012345.0
    return str(v).strip()


def read_rows(data: bytes, filename: str) -> list[list[str]]:
    fn = filename.lower()
    try:
        if fn.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            rows = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        elif fn.endswith(".xls"):
            import xlrd
            ws = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
            rows = [[_cell(ws.cell_value(r, c)) for c in range(ws.ncols)] for r in range(ws.nrows)]
        elif fn.endswith((".csv", ".txt")):
            text = data.decode("utf-8-sig", errors="replace")
            rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
        else:
            raise SheetError("Unsupported sheet type. Upload .xlsx, .xls or .csv.")
    except SheetError:
        raise
    except Exception as e:  # noqa: BLE001
        raise SheetError(f"Could not open the sheet ({type(e).__name__}). Check it is a valid Excel/CSV file.") from e
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise SheetError("The sheet is empty.")
    return rows


def _map_header(row: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, h in enumerate(row):
        n = re.sub(r"[^a-z0-9]", "", h.lower())
        hit = process.extractOne(n, _FLAT.keys(), scorer=fuzz.ratio, score_cutoff=85) if n else None
        if hit and _FLAT[hit[0]] not in out:
            out[_FLAT[hit[0]]] = i
    return out


def _map_by_content(rows: list[list[str]]) -> dict[str, int]:
    width = max(len(r) for r in rows)
    out: dict[str, int] = {}
    for i in range(width):
        col = [r[i] for r in rows if i < len(r) and r[i]]
        if not col:
            continue
        if sum("@" in c for c in col) >= len(col) / 2 and "email" not in out:
            out["email"] = i
        elif sum(len(re.sub(r"\D", "", c)) >= 9 for c in col) >= len(col) / 2 and "phone" not in out:
            out["phone"] = i
        elif "name" not in out:
            out["name"] = i
    return out


def parse_sheet(data: bytes, filename: str) -> tuple[list[dict], list[dict]]:
    """Returns (accepted rows, rejected rows). Each accepted row: {name, email, phone(10 digits)}."""
    rows = read_rows(data, filename)
    mapping = _map_header(rows[0])
    if len(mapping) >= 2:
        body = rows[1:]
    else:
        mapping, body = _map_by_content(rows), rows
    if len(mapping) < 3:
        raise SheetError("I need three columns: name, email and phone number. "
                         f"I could only identify: {', '.join(mapping) or 'none'}.")
    ok, bad, seen = [], [], set()
    for n, r in enumerate(body, start=2 if body is not rows else 1):
        g = lambda k: r[mapping[k]].strip() if mapping[k] < len(r) else ""  # noqa: E731
        name, email, phone = re.sub(r"\s+", " ", g("name")), g("email"), g("phone")
        if not (name or email or phone):
            continue
        reasons = []
        if not name:
            reasons.append("name is missing")
        if not valid_email(email):
            reasons.append("email is not valid")
        p = normalize_phone(phone)
        if not p:
            reasons.append("phone is not a valid 10-digit mobile")
        key_e, key_p = email.lower(), p
        if not reasons and (key_e in seen or key_p in seen):
            reasons.append("duplicate of an earlier row in this sheet")
        if reasons:
            bad.append({"row": n, "name": name or "(no name)", "reason": ", ".join(reasons)})
            continue
        seen.update({key_e, key_p})
        ok.append({"name": name, "email": email, "phone": p})
    return ok, bad
