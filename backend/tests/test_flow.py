"""End to end through the HTTP API, exactly as the chat window uses it. No sample files: everything is built here."""
import io
import json
import uuid

import pytest
import xlrd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app, cfg
from tests.conftest import PNG, make_xls, make_xlsx

BATCH = [["Manohar Lal", "mano.lal@gmail.com", "9845012345"], ["Asha Rao", "asha.rao@company.com", 9845022222],
         ["Varun Thakkar", "varun22@hotmail.com", "+91 81234 56789"], ["Priyadarshini Mani", "priya.mani@yahoo.com.in", "9900011111"]]


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


class Chatter:
    def __init__(self, client):
        self.c, self.sid = client, uuid.uuid4().hex

    def say(self, text="", files=()):
        r = self.c.post("/api/chat", data={"session_id": self.sid, "text": text}, files=[("files", f) for f in files])
        assert r.status_code == 200, r.text
        self.last = r.json()
        return "\n".join(m["text"] for m in self.last["messages"])

    def card(self, kind):
        return next((m["card"] for m in self.last["messages"] if m["card"] and m["card"]["kind"] == kind), None)


IMG = ("photo.png", PNG, "image/png")


def test_api_basics(client):
    assert client.get("/api/health").json()["employees"] == 0
    assert {"employee_id", "full_name", "age", "phone", "address", "aadhaar", "pan", "bank", "documents", "verification"} <= client.get("/api/schema").json().keys()
    assert client.post("/api/chat", data={"session_id": "x", "text": "hi"}).status_code == 400
    assert client.get("/api/chat/bad id!").status_code in (400, 404)
    assert client.get("/api/master/download").status_code == 404
    assert client.get("/api/master/download?format=pdf").status_code == 400


def test_before_a_sheet(client):
    c = Chatter(client)
    assert "no employee waiting" in c.say("", [IMG])
    assert "Upload an employee sheet" in c.say("hello") and c.last["chips"] == ["Help"]
    assert "can't use" in c.say("", [("notes.docx", b"x" * 900, "application/octet-stream")])
    assert "need three columns" in c.say("", [("bad.csv", b"Name,Age\nA,3\n", "text/csv")])
    assert "Anytime" in c.say("thanks")
    assert "didn't quite get that" in c.say("blah blah blah")


def test_whole_journey(client):
    c = Chatter(client)
    out = c.say("Onboard these", [("batch.xlsx", make_xlsx(BATCH), "application/octet-stream")])
    assert "4 accepted" in out and "Employee 1 of 4: Manohar Lal" in out and "Send image 1 of 4: **PAN card**" in out
    assert c.last["progress"] == {"phase": "collecting", "total": 4, "index": 1, "name": "Manohar Lal", "done": 0}

    # four images, one step each: PAN, Aadhaar, cheque, any other. An image fills ONLY its own step.
    out = c.say("", [IMG])
    assert out.startswith("Yes ✓ PAN card filled") and "Send image 2 of 4: **Aadhaar card**" in out
    e = c.c.get("/api/chat/" + c.sid).json()                           # nothing saved to the master list yet
    assert client.get("/api/master").json()["employees"] == []
    out = c.say("", [IMG])
    assert "Aadhaar card filled" in out and "XXXX XXXX" in out and "Send image 3 of 4: **cancelled cheque**" in out
    out = c.say("", [IMG])
    assert "Cancelled cheque filled" in out and "Send image 4 of 4: **any other document**" in out
    out = c.say("", [IMG])
    assert "Any other document filled" in out and "EMP-0001" in out and "Employee 2 of 4: Asha Rao" in out and "image 1 of 4" in out
    m = client.get("/api/master").json()["employees"][0]
    assert m["pan"]["number"] and m["aadhaar"]["number_masked"] and m["bank"]["ifsc"] and m["other"]["reference"].startswith("OTH-")
    assert all(d["source"] == "placeholder" and len(d["files"]) == 1 for d in m["documents"].values())

    # smart questions mid-way
    out = c.say("how many are left?")
    assert "1 of 4" in out and "Now on **Asha Rao**" in out and "next up: Varun Thakkar" in out
    assert "Manohar Lal" in c.say("show manohar") and c.card("fields")["title"] == "Manohar Lal"
    assert "Asha Rao" in c.say("what about the second one?")
    assert "missing" in c.say("what's missing?").lower()

    # plain-language edits (short replies), a pronoun after naming her
    out = c.say("change Asha's phone to 9845099999")
    assert out == "Done: phone set to 9845099999."
    assert "email set to asha@new.com" in c.say("her email is asha@new.com")
    assert "valid" in c.say("fix phone to 123")
    # several files in ONE message are used one step each, in order (PDF and other image types are fine)
    out = c.say("", [("scan.pdf", b"%PDF-1.4 anything", "application/pdf"), IMG, ("a.jpg", PNG, "image/jpeg"), IMG])
    assert out.count("filled") == 4 and "EMP-0002" in out and "Employee 3 of 4: Varun Thakkar" in out

    # real details typed in are kept; the matching image fills only what is still empty
    out = c.say("PAN ABCDE1234F, ifsc SBIN0001234 and account 123456789012")
    assert "pan number set to ABCDE1234F" in out and "Send image 1 of 4: **PAN card**" in out
    for _ in range(4):
        out = c.say("", [("a.jpg", PNG, "image/jpeg")])
    assert "EMP-0003" in out and "Employee 4 of 4" in out
    v = next(e for e in client.get("/api/master").json()["employees"] if e["full_name"] == "Varun Thakkar")
    assert v["pan"]["number"] == "ABCDE1234F" and v["bank"]["ifsc"] == "SBIN0001234" and v["bank"]["account_number"] == "123456789012"
    assert v["aadhaar"]["number_masked"].startswith("XXXX") and v["pan"]["name"] == "VARUN THAKKAR"

    # skip the last one; the batch goes to review automatically
    out = c.say("skip this one")
    assert "Skipped **Priyadarshini Mani**" in out and "3 of 3" in out and c.last["progress"]["phase"] == "review"
    rows = c.card("report")["rows"]
    assert [r["result"] for r in rows] == ["PASS", "PASS", "PASS", "SKIPPED"] and "Placeholder values" in " ".join(rows[0]["warnings"])
    assert c.last["chips"] == ["Confirm", "What's wrong?", "Status"]

    # the master JSON is the standard record; the Aadhaar number is never stored in full
    emps = client.get("/api/master").json()["employees"]
    assert len(emps) == 3 and set(emps[0]) == set(client.get("/api/schema").json())
    m = emps[0]
    assert m["employee_id"] == "EMP-0001" and m["dob"] == "01/01/1990" and m["age"] and "number" not in m["aadhaar"] and len(m["aadhaar"]["hash"]) == 64
    assert m["pan"]["name"] == "MANOHAR LAL" and m["documents"]["pan"]["files"]

    # confirm in words ("looks good"), then download both formats
    assert "Published **3**" in c.say("looks good") and c.card("download")
    assert [e["status"] for e in client.get("/api/master").json()["employees"]] == ["VERIFIED"] * 3
    rows = list(load_workbook(io.BytesIO(client.get("/api/master/download?format=xlsx").content)).active.iter_rows(values_only=True))
    head = list(rows[0])
    assert [r[1] for r in rows[1:]] == ["Manohar Lal", "Asha Rao", "Varun Thakkar"] and rows[2][head.index("Phone")] == "9845099999"
    assert rows[1][head.index("Data")] == "placeholder" and rows[1][head.index("Aadhaar (masked)")].startswith("XXXX")
    xls = xlrd.open_workbook(file_contents=client.get("/api/master/download?format=xls").content).sheet_by_index(0)
    assert xls.nrows == 4 and xls.cell_value(0, 1) == "Full Name"
    assert "Here's the master sheet" in c.say("download the excel please")

    # editing a published employee withdraws them from the sheet until they are verified again
    assert "pan number set to PQRSP9876K" in c.say("change EMP-0001's pan to PQRSP9876K")
    assert "EMP-0001" in c.say("show EMP-0001") and "PQRSP9876K" in json.dumps(c.card("fields"))
    assert "Done" in c.say("set Manohar's name to Manohar Lal Sharma")
    assert [e["status"] for e in client.get("/api/master").json()["employees"]] == ["COMPLETE", "VERIFIED", "VERIFIED"]
    assert len(list(load_workbook(io.BytesIO(client.get("/api/master/download").content)).active.iter_rows())) == 3
    out = c.say("what's wrong with manohar?")                       # a name in a sentence, not a number
    assert "issue" in out and "does not match the sheet name" in out
    assert "No problems" in c.say("any problems with asha?")
    assert "Done" in c.say("set manohar's name to Manohar Lal")
    assert "1 of 3 passed" not in c.say("verfiy") and c.last["chips"][0] == "Confirm"      # typo'd command still understood
    assert "Published **3**" in c.say("confrim")

    # duplicates: rejected at the door, caught when details are typed, and when editing
    d = Chatter(client)
    out = d.say("", [("b2.xlsx", make_xlsx([["Manohar Lal", "mano.lal@gmail.com", "9845012345"], ["Ravi Kumar", "ravi@company.com", "9988776655"],
                                              ["Bad Email", "nope", "9845033333"]]), "x")])
    assert "3 rows, **1 accepted**, 2 skipped" in out
    reasons = {r[1]: r[2] for r in d.card("skipped")["rows"]}
    assert "already in the master list as EMP-0001" in reasons["Manohar Lal"] and "email is not valid" in reasons["Bad Email"]
    out = d.say("PAN " + v["pan"]["number"])
    assert "Duplicate" in out and "EMP-0003" in out and "was not added" in out
    assert len(client.get("/api/master").json()["employees"]) == 3
    assert "Not changed" in c.say(f"change Asha's pan to {v['pan']['number']}")

    # a restart changes nothing
    from app.chat import Chat
    from app.sessions import Sessions
    from app.store import MasterStore
    st2 = MasterStore(cfg.data_dir)
    assert len(st2.all()) == 3
    h = Chat(st2, Sessions(cfg.data_dir), cfg).history(c.sid)
    assert h["messages"][0]["role"] == "user" and h["progress"]["phase"] == "review"


def test_xls_input_reset_and_ambiguity(client):
    c = Chatter(client)
    rows = [["Anil Kumar", "anil.k@x.com", "9811100001"], ["Anil Kumari", "anil.ku@x.com", "9811100002"], ["Zoya Khan", "zoya@x.com", "9811100003"]]
    out = c.say("", [("s.xls", make_xls(rows), "x")])
    assert "3 accepted" in out
    assert "Which one do you mean" in c.say("show anil")
    assert "Zoya Khan" in c.say("show zoya") and "Zoya Khan" in c.say("tell me about #3")
    assert "A batch is still in progress" in c.say("", [("s.xls", make_xls(rows), "x")])
    assert "Cleared" in c.say("start over") and c.last["progress"]["phase"] == "idle"
    assert "Nothing to skip" in c.say("skip")
    assert "I couldn't tell who" in c.say("change phone to 9811100009")
