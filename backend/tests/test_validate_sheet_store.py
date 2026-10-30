import io
import json
import tempfile
from datetime import date

import pytest

from app import sheet, store as st, validate as V
from app.schema import TEMPLATE, new_employee
from tests.conftest import make_xls, make_xlsx


# ---------------- validators ----------------
def test_ids():
    assert V.valid_pan("ABCPL1234F") and not V.valid_pan("ABCPL12345") and not V.valid_pan("ABC1234567")
    assert V.valid_ifsc("SBIN0001234") and not V.valid_ifsc("SBIN1001234")
    assert V.valid_account("123456789012") and not V.valid_account("12345")
    base = "23456789012"
    num = base + V.verhoeff_digit(base)
    assert V.valid_aadhaar(num) and not V.valid_aadhaar(num[:-1] + str((int(num[-1]) + 1) % 10)) and not V.valid_aadhaar("123456789012")
    assert V.mask_aadhaar("2345 6789 0124") == "XXXX XXXX 0124"


def test_phone_email_dates():
    assert V.normalize_phone("98450 12345") == "9845012345" and V.normalize_phone("+91 81234-56789") == "8123456789"
    assert V.normalize_phone("12345") == "" and V.normalize_phone("5845012345") == "" and V.normalize_phone("9xxxx12345") == ""
    assert V.valid_email("a.b@x.co.in") and not V.valid_email("nope") and not V.valid_email("a@b")
    assert V.calc_age("15/08/1990", date(2026, 10, 12)) == 36 and V.calc_age("15/11/1990", date(2026, 10, 12)) == 35
    assert V.calc_age("1990", date(2026, 10, 12)) == 36 and V.calc_age("") is None


def test_name_matching():
    assert V.name_status("Manohar Lal", "MANOHARLAL")[0] == "PASS"          # OCR glued the words
    assert V.name_status("Lal Manohar", "MANOHAR LAL")[0] == "PASS"         # order
    assert V.name_status("Mr. Rahul Kumar", "RAHUL KUMAR")[0] == "PASS"     # title
    assert V.name_status("Monohar Lal", "MANOHAR LAL")[0] == "WARN"         # one letter off
    assert V.name_status("Priyadarshini Mani", "PRIYAM")[0] == "FAIL"
    assert V.name_status("", "X")[0] == "MISSING"


def _complete():
    e = new_employee("Asha Rao", "asha@x.com", "9845022222")
    e["aadhaar"].update(number_masked="XXXX XXXX 1238", checksum_valid=True, name="ASHA RAO", dob="09/04/1995")
    e["pan"].update(number="BCDPR5678G", name="ASHA RAO", father_name="K RAO", dob="09/04/1995")
    e["bank"].update(account_holder="ASHA RAO", account_number="987654321098", ifsc="HDFC0004321")
    e["address"]["full"] = "7 Lake View, Bengaluru"
    e["other"]["reference"] = "OTH-123456"
    e["dob"], e["age"] = "09/04/1995", V.calc_age("09/04/1995")
    return e


def test_verify_rules():
    assert V.verify(_complete())["passed"]
    e = _complete(); e["pan"]["number"] = "BAD"; e["aadhaar"]["checksum_valid"] = False; e["bank"]["account_holder"] = "SOMEONE ELSE"
    e["pan"]["dob"] = "01/01/1990"; e["address"]["full"] = ""
    issues = " | ".join(V.verify(e)["issues"])
    for needle in ("PAN BAD", "checksum", "bank account holder", "Date of birth differs", "Missing: address"):
        assert needle in issues
    e = _complete(); e["aadhaar"]["dob"] = e["dob"] = "01/01/2015"; e["pan"]["dob"] = "01/01/2015"; e["age"] = V.calc_age(e["dob"])
    assert any("at least 18" in i for i in V.verify(e)["issues"])
    assert any("Duplicate PAN" in i for i in V.verify(_complete(), [{"field": "PAN", "employee_id": "EMP-0009", "name": "X"}])["issues"])


def test_template_file_matches_code():
    from pathlib import Path
    on_disk = json.loads((Path(__file__).resolve().parents[1] / "schema" / "employee_master.template.json").read_text())
    assert on_disk == TEMPLATE


# ---------------- sheets ----------------
ROWS = [["Manohar Lal", "mano.lal@gmail.com", "98450 12345"], ["Asha Rao", "asha.rao@company.com", 9845022222],
        ["Varun Thakkar", "varun22@hotmail.com", "+91 81234 56789"], ["Priyadarshini  Mani", "priya.mani@yahoo.com.in", "9900011111"]]


def test_sheet_xlsx_xls_csv():
    ok, bad = sheet.parse_sheet(make_xlsx(ROWS), "a.xlsx")
    assert not bad and [r["name"] for r in ok] == ["Manohar Lal", "Asha Rao", "Varun Thakkar", "Priyadarshini Mani"]  # double space collapsed
    assert ok[0]["phone"] == "9845012345" and ok[1]["phone"] == "9845022222" and ok[2]["phone"] == "8123456789"     # numeric cell, +91 format
    ok, bad = sheet.parse_sheet(make_xls([["Manohar Lal", "mano.lal@gmail.com", 9845012345.0]]), "a.xls")        # .xls with a float phone
    assert ok == [{"name": "Manohar Lal", "email": "mano.lal@gmail.com", "phone": "9845012345"}]
    csv = b"Employee Name,Email ID,Mobile\nAsha Rao,asha.rao@company.com,9845022222\n"
    assert sheet.parse_sheet(csv, "a.csv")[0][0]["name"] == "Asha Rao"


def test_sheet_no_header_and_problems():
    ok, bad = sheet.parse_sheet(make_xlsx([["Neha Verma", "neha@company.com", "9811122233"]], header=None), "a.xlsx")
    assert ok == [{"name": "Neha Verma", "email": "neha@company.com", "phone": "9811122233"}] and not bad
    rows = [["Manohar Lal", "mano.lal@gmail.com", "9845012345"], ["Ravi Kumar", "ravi.kumar@company.com", "9988776655"], ["Bad Email", "not-an-email", "9845033333"],
            ["No Phone", "no.phone@x.com", ""], ["Short Number", "short@x.com", "12345"], ["Dup In Sheet", "ravi.kumar@company.com", "9000011111"]]
    ok, bad = sheet.parse_sheet(make_xlsx(rows), "a.xlsx")
    reasons = {b["name"]: b["reason"] for b in bad}
    assert [r["name"] for r in ok] == ["Manohar Lal", "Ravi Kumar"]
    assert "email is not valid" in reasons["Bad Email"] and "phone" in reasons["No Phone"] and "phone" in reasons["Short Number"]
    assert "duplicate of an earlier row" in reasons["Dup In Sheet"]


def test_sheet_errors():
    with pytest.raises(sheet.SheetError):
        sheet.parse_sheet(b"", "a.csv")
    with pytest.raises(sheet.SheetError):
        sheet.parse_sheet(b"Name,Age\nA,3\n", "a.csv")
    with pytest.raises(sheet.SheetError):
        sheet.parse_sheet(b"junk", "a.xlsx")
    with pytest.raises(sheet.SheetError):
        sheet.parse_sheet(b"x", "a.pdf")


# ---------------- store ----------------
def _emp(n, **kw):
    e = new_employee(f"Person {n}", f"p{n}@x.com", f"98450000{n:02d}")
    for k, v in kw.items():
        a, b = k.split("__")
        e[a][b] = v
    return e


def test_store_ids_duplicates_and_sheet():
    s = st.MasterStore(tempfile.mkdtemp())
    a = s.add(_emp(1, pan__number="ABCPL1234F", aadhaar__hash="h1", bank__account_number="111111111", bank__ifsc="SBIN0001234"))
    assert a["employee_id"] == "EMP-0001" and a["status"] == "COMPLETE"
    for kw, field in (({}, "email"), ({"pan__number": "ABCPL1234F"}, "PAN"), ({"aadhaar__hash": "h1"}, "Aadhaar"),
                      ({"bank__account_number": "111111111", "bank__ifsc": "SBIN0009999"}, "bank account")):
        e = _emp(2, **kw)
        if field == "email":
            e["email"] = "p1@x.com"
        with pytest.raises(st.Duplicate) as ei:
            s.add(e)
        assert ei.value.conflicts[0]["field"] == field and ei.value.conflicts[0]["employee_id"] == "EMP-0001"
    e = _emp(3); e["phone"] = "9845000001"
    with pytest.raises(st.Duplicate):
        s.add(e)
    b = s.add(_emp(4))
    assert b["employee_id"] == "EMP-0002" and len(s.all()) == 2
    # editing an employee into someone else's PAN is blocked; its own values are fine
    b["pan"]["number"] = "ABCPL1234F"
    with pytest.raises(st.Duplicate):
        s.update(b)
    b["pan"]["number"] = "XYZPL9999Z"
    s.update(b)
    # master sheet holds only VERIFIED rows and is rebuilt (never appended), so re-publishing cannot duplicate
    assert s.publish_excel() == 0
    for e in s.all():
        e["status"] = "VERIFIED"; e["verification"]["checked_at"] = "2026-10-12T10:00:00Z"
        s.update(e, publish=False)
    assert s.publish_excel() == 2 and s.publish_excel() == 2
    from openpyxl import load_workbook
    import xlrd
    ws = load_workbook(s.xlsx_path).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Employee ID", "Full Name", "Email", "Phone") and [r[0] for r in rows[1:]] == ["EMP-0001", "EMP-0002"]
    xs = xlrd.open_workbook(str(s.xls_path)).sheet_by_index(0)
    assert xs.nrows == 3 and xs.cell_value(1, 0) == "EMP-0001" and xs.cell_value(1, 3) == "9845000001"


def test_store_persistence_and_corruption_recovery():
    d = tempfile.mkdtemp()
    s = st.MasterStore(d)
    s.add(_emp(1)); s.add(_emp(2))
    assert len(st.MasterStore(d).all()) == 2                      # survives a restart
    s.json_path.write_text("{ not json")                          # simulate a torn/corrupt file
    assert len(st.MasterStore(d).all()) == 1                      # falls back to the .bak: the state before the last write
